"""Excel Sink — 写入 Excel 文件"""

from __future__ import annotations

from pathlib import Path

from openspider.core.sinks.base import BaseSink


class ExcelSink(BaseSink):
    """将爬取数据写入 Excel (.xlsx) 文件"""

    def __init__(self, config: dict):
        super().__init__(config)
        self.file_path = Path(config.get("path", f"./data/{self._table_name()}.xlsx"))
        self._rows: list[dict] = []

    async def open(self):
        self.file_path.parent.mkdir(parents=True, exist_ok=True)
        # 加载已有数据
        if self.file_path.exists():
            try:
                from openpyxl import load_workbook
                wb = load_workbook(self.file_path)
                ws = wb.active
                headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    self._rows.append(dict(zip(headers, row)))
                wb.close()
            except Exception:
                pass

    async def close(self):
        await self.flush()
        self._save()

    async def _write_batch(self, items: list[dict]):
        self._rows.extend(items)

    def _save(self):
        if not self._rows:
            return
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = self.spider_name[:31]  # Excel sheet 名最长 31 字符
            # 写表头
            headers = list(self._rows[0].keys())
            ws.append(headers)
            # 写数据
            for row in self._rows:
                ws.append([row.get(h, "") for h in headers])
            wb.save(self.file_path)
            wb.close()
        except ImportError:
            # openpyxl 未安装，回退为 CSV
            import csv
            fallback = self.file_path.with_suffix(".csv")
            with open(fallback, "w", newline="", encoding="utf-8") as f:
                if self._rows:
                    writer = csv.DictWriter(f, fieldnames=self._rows[0].keys())
                    writer.writeheader()
                    writer.writerows(self._rows)
